import csv
from django.db.models import Q
from django.http import HttpResponse, JsonResponse
from rest_framework import status
from rest_framework.decorators import api_view, authentication_classes, permission_classes
from rest_framework.permissions import AllowAny
from rest_framework.response import Response

from .models import OlympiadEvent, Participant, Certificate, SiteMediaSettings
from .serializers import (
    ParticipantSerializer,
    EventInfoSerializer,
    CertificateSerializer,
    SiteMediaSettingsSerializer
)


@api_view(['POST'])
@authentication_classes([])
@permission_classes([AllowAny])
def register_participant(request):
    """
    POST /api/v1/participants/
    Register a new participant for the currently active olympiad event.
    """
    # Find the active event
    event = OlympiadEvent.objects.filter(is_active=True).first()
    if not event:
        # If no event exists yet, create or get a default active event
        event, _ = OlympiadEvent.objects.get_or_create(
            slug='viloyat-matematika-olimpiadasi-2026',
            defaults={
                'title': 'Viloyat matematika olimpiadasi',
                'is_active': True,
                'location': "Farg'ona shahri",
            }
        )

    serializer = ParticipantSerializer(
        data=request.data,
        context={'event': event}
    )

    if serializer.is_valid():
        participant = serializer.save(event=event)
        return Response(
            {
                'id': participant.id,
                'message': "Ro'yxatdan o'tish muvaffaqiyatli yakunlandi",
            },
            status=status.HTTP_201_CREATED
        )

    return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


@api_view(['GET'])
def list_participants(request):
    """
    GET /api/v1/applications/
    Return list of participants for the dashboard with search, filters, and stats.
    """
    queryset = Participant.objects.all().order_by('-created_at')

    search = request.GET.get('search', '').strip()
    if search:
        queryset = queryset.filter(
            Q(full_name__icontains=search) |
            Q(phone__icontains=search)
        )

    class_number = request.GET.get('class_number')
    if class_number and class_number.isdigit():
        queryset = queryset.filter(class_number=int(class_number))

    serializer = ParticipantSerializer(queryset, many=True)

    all_participants = Participant.objects.all()
    stats = {
        'total': all_participants.count(),
        'class_5': all_participants.filter(class_number=5).count(),
        'class_6': all_participants.filter(class_number=6).count(),
        'class_7': all_participants.filter(class_number=7).count(),
    }

    return Response({
        'stats': stats,
        'count': queryset.count(),
        'results': serializer.data
    })


@api_view(['DELETE'])
def delete_participant(request, pk):
    """
    DELETE /api/v1/participants/<pk>/
    Delete a participant registration from dashboard.
    """
    try:
        participant = Participant.objects.get(pk=pk)
        participant.delete()
        return Response({'message': "Muvaffaqiyatli o'chirildi"}, status=status.HTTP_200_OK)
    except Participant.DoesNotExist:
        return Response({'error': "Ishtirokchi topilmadi"}, status=status.HTTP_404_NOT_FOUND)


@api_view(['GET'])
def export_participants_excel(request):
    """
    GET /api/v1/applications/export/
    Export participants list to a styled Excel (.xlsx) file.
    """
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    queryset = Participant.objects.all().order_by('-created_at')

    search = request.GET.get('search', '').strip()
    if search:
        queryset = queryset.filter(
            Q(full_name__icontains=search) |
            Q(phone__icontains=search)
        )

    class_number = request.GET.get('class_number')
    if class_number and class_number.isdigit():
        queryset = queryset.filter(class_number=int(class_number))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Qabullar"

    # Header styling
    header_fill = PatternFill(start_color="1A2035", end_color="1A2035", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style='thin', color='D3D3D3'),
        right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'),
        bottom=Side(style='thin', color='D3D3D3')
    )

    headers = ['№ ID', 'F.I.SH. (O\'quvchi)', 'Telefon raqami', 'Sinfi', 'Ro\'yxatdan o\'tgan vaqti']
    ws.append(headers)
    ws.row_dimensions[1].height = 26

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border

    # Data rows
    row_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    for row_idx, p in enumerate(queryset, start=2):
        created_str = p.created_at.strftime('%Y-%m-%d %H:%M') if p.created_at else ''
        ws.append([
            p.id,
            p.full_name,
            p.phone,
            f"{p.class_number}-sinf",
            created_str
        ])
        ws.row_dimensions[row_idx].height = 20

        # Style cells
        ws.cell(row=row_idx, column=1).alignment = center_align
        ws.cell(row=row_idx, column=2).alignment = left_align
        ws.cell(row=row_idx, column=3).alignment = center_align
        ws.cell(row=row_idx, column=4).alignment = center_align
        ws.cell(row=row_idx, column=5).alignment = center_align

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = row_font
            cell.border = thin_border
            if row_idx % 2 == 1:
                cell.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    # Column widths auto-fit
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="qabullar_olimpiada.xlsx"'
    return response


@api_view(['GET'])
def event_info(request):
    """
    GET /api/v1/event-info/
    Return public information about the currently active olympiad.
    """
    event = OlympiadEvent.objects.filter(is_active=True).first()
    if not event:
        event, _ = OlympiadEvent.objects.get_or_create(
            slug='viloyat-matematika-olimpiadasi-2026',
            defaults={
                'title': 'Viloyat matematika olimpiadasi',
                'is_active': True,
                'location': "Farg'ona shahri",
            }
        )
    serializer = EventInfoSerializer(event)
    data = serializer.data
    data['total_registered'] = Participant.objects.count()
    return Response(data)


@api_view(['GET'])
def list_certificates(request):
    """
    GET /api/v1/certificates/
    Return active certificates / gallery items for the frontend carousel.
    """
    certificates = Certificate.objects.filter(is_active=True).order_by('order', '-created_at')
    serializer = CertificateSerializer(certificates, many=True, context={'request': request})
    return Response(serializer.data)


@api_view(['GET'])
def site_media_settings(request):
    """
    GET /api/v1/site-settings/
    Return dynamic logo and site images uploaded from Django admin.
    """
    settings = SiteMediaSettings.get_settings()
    serializer = SiteMediaSettingsSerializer(settings, context={'request': request})
    return Response(serializer.data)


def health_check(request):
    """Simple health check endpoint."""
    return JsonResponse({'status': 'ok'})


